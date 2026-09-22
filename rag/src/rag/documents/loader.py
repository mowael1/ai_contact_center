"""Format dispatch for document loading."""

from __future__ import annotations

from pathlib import Path

from rag.documents.base import HeadingCandidate, LoadedDocument, PageText, infer_title
from rag.logging_utils import get_logger

logger = get_logger(__name__)

SUPPORTED = {".pdf", ".txt", ".md", ".markdown", ".docx"}


def load_document(path: str | Path) -> LoadedDocument:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from rag.documents.pdf_loader import load_pdf

        return load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix in (".txt", ".md", ".markdown"):
        return _load_text(path)
    raise ValueError(f"Unsupported file type {suffix!r}. Supported: {sorted(SUPPORTED)}")


def iter_supported_files(directory: str | Path) -> list[Path]:
    directory = Path(directory)
    if not directory.exists():
        raise FileNotFoundError(f"Directory not found: {directory}")
    return sorted(
        p for p in directory.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith(".")
    )


def _load_text(path: Path) -> LoadedDocument:
    """Plain text / Markdown. Markdown ``#`` levels are real headings."""
    import re

    raw = path.read_text(encoding="utf-8", errors="replace")
    headings: list[HeadingCandidate] = []
    for order, line in enumerate(raw.splitlines()):
        match = re.match(r"^(#{1,6})\s+(.*\S)\s*$", line)
        if match:
            headings.append(
                HeadingCandidate(text=match.group(2), page_number=1,
                                 level=len(match.group(1)), order=order,
                                 reason="markdown")
            )
    return LoadedDocument(
        title=infer_title(path, raw[:400]),
        source=path.name,
        path=str(path.resolve()),
        source_type="md" if path.suffix.lower() != ".txt" else "txt",
        pages=[PageText(page_number=1, text=raw.strip())],
        headings=headings,
        byte_size=path.stat().st_size,
    )


def _load_docx(path: Path) -> LoadedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required for .docx files: pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    lines: list[str] = []
    headings: list[HeadingCandidate] = []
    for order, para in enumerate(document.paragraphs):
        text = (para.text or "").strip()
        if not text:
            continue
        lines.append(text)
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            try:
                level = int(style.split()[-1])
            except (ValueError, IndexError):
                level = 1
            headings.append(
                HeadingCandidate(text=text, page_number=1, level=level,
                                 order=order, reason="docx_style")
            )
    return LoadedDocument(
        title=infer_title(path, "\n".join(lines[:5])),
        source=path.name,
        path=str(path.resolve()),
        source_type="docx",
        pages=[PageText(page_number=1, text="\n".join(lines))],
        headings=headings,
        byte_size=path.stat().st_size,
    )
